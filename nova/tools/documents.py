"""Reading documents and pictures that are not plain text.

extract_text turns PDF, Word, Excel and PowerPoint files - and on macOS the
older formats textutil understands - into text for read_file, which numbers and
pages it like any other file. view_image hands a picture file to the vision
model. Each format's library is imported only when that format is read, so a
missing one fails just that format, with a message saying what to install.
"""
from __future__ import annotations

import functools
import importlib.util
import subprocess
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

from ..attachments import DOCUMENT_KINDS, IMAGE_SUFFIXES
from ..config import Config
from ..llm import LLM
from .registry import Registry, schema, string

MAX_SHEET_ROWS = 5_000
DEFAULT_QUESTION = "Describe this image in detail, including any text in it."


class DocumentError(Exception):
    """A document could not be read; the message is fit to hand to the model."""


@dataclass(frozen=True)
class Extracted:
    text: str
    summary: str  # for read_file's header, e.g. "PDF, 12 pages"


def is_document(path: Path) -> bool:
    return path.suffix.lower() in DOCUMENT_KINDS


def extract_text(path: Path) -> Extracted:
    """The readable text of a document. Raises DocumentError, or
    PermissionError when the operating system refuses access."""
    stat = path.stat()
    return _extract(str(path), stat.st_mtime_ns, stat.st_size)


@functools.lru_cache(maxsize=4)
def _extract(path: str, _mtime_ns: int, _size: int) -> Extracted:
    # The modification time and size are only there to turn an edited file into
    # a cache miss, so paging through a long PDF reads it once.
    target = Path(path)
    suffix = target.suffix.lower()
    try:
        return _READERS[suffix](target)
    except (DocumentError, PermissionError):
        raise
    except ImportError as exc:
        raise DocumentError(
            f"reading {suffix} needs the {_PACKAGES.get(suffix, 'right')} package."
        ) from exc
    except Exception as exc:
        raise DocumentError(f"could not read {target.name}: {exc}") from exc


def _plural(count: int, noun: str) -> str:
    return f"{count} {noun}{'' if count == 1 else 's'}"


def _pdf(path: Path) -> Extracted:
    from pypdf import PdfReader

    reader = PdfReader(path)
    # Many "protected" PDFs only restrict printing and open with no password.
    if reader.is_encrypted and not reader.decrypt(""):
        raise DocumentError(f"{path.name} is password-protected.")
    pages = [page.extract_text() or "" for page in reader.pages]
    if not any(text.strip() for text in pages):
        raise DocumentError(
            f"no text found in {path.name}; it looks like a scan, which read_file "
            "cannot read."
        )
    body = "\n".join(f"--- page {n} ---\n{text.strip()}" for n, text in enumerate(pages, 1))
    return Extracted(body, f"PDF, {_plural(len(pages), 'page')}")


def _docx(path: Path) -> Extracted:
    """Paragraphs and tables in document order, read straight from the XML.

    python-docx's own paragraph text skips fill-in fields (content controls)
    and anything they wrap, which are exactly where a contract or form keeps
    the parts that matter.
    """
    import docx
    from docx.oxml.ns import qn

    fallback = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"

    def text_of(element) -> str:
        parts: list[str] = []

        def walk(node) -> None:
            for child in node:
                if child.tag == fallback:  # the same content again, for old Word
                    continue
                if child.tag == qn("w:t"):
                    parts.append(child.text or "")
                elif child.tag in (qn("w:tab"), qn("w:br")) and node.tag == qn("w:r"):
                    parts.append("\t" if child.tag == qn("w:tab") else "\n")
                else:
                    walk(child)

        walk(element)
        return "".join(parts)

    def blocks(parent):
        for child in parent.iterchildren():
            if child.tag == qn("w:p"):
                yield text_of(child)
            elif child.tag == qn("w:tbl"):
                for row in child.iterchildren(qn("w:tr")):
                    cells = row.iterchildren(qn("w:tc"))
                    yield " | ".join(" ".join(blocks(cell)).strip() for cell in cells)
            elif child.tag == qn("w:sdt"):  # a fill-in field around whole paragraphs
                content = child.find(qn("w:sdtContent"))
                if content is not None:
                    yield from blocks(content)

    body = docx.Document(str(path)).element.body
    return Extracted("\n".join(blocks(body)), "Word document")


def _xlsx(path: Path) -> Extracted:
    import openpyxl

    # data_only: the values Excel last calculated, not the formulas behind them.
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        lines = []
        for sheet in workbook.worksheets:
            lines.append(f"=== Sheet: {sheet.title} ===")
            rows = 0
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if value is None else str(value) for value in row]
                while cells and not cells[-1]:
                    cells.pop()
                if not cells:
                    continue
                if rows == MAX_SHEET_ROWS:
                    lines.append(f"... stopped at {MAX_SHEET_ROWS:,} rows; the sheet has more")
                    break
                lines.append(" | ".join(cells))
                rows += 1
        sheets = len(workbook.worksheets)
    finally:
        workbook.close()
    return Extracted("\n".join(lines), f"Excel workbook, {_plural(sheets, 'sheet')}")


def _pptx(path: Path) -> Extracted:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def shape_text(shapes):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from shape_text(shape.shapes)
            elif shape.has_table:
                for row in shape.table.rows:
                    yield " | ".join(cell.text.strip() for cell in row.cells)
            elif shape.has_text_frame and shape.text_frame.text.strip():
                yield shape.text_frame.text.strip()

    deck = Presentation(str(path))
    lines = []
    for number, slide in enumerate(deck.slides, 1):
        lines.append(f"--- slide {number} ---")
        lines.extend(shape_text(slide.shapes))
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            notes = frame.text.strip() if frame is not None else ""
            if notes:
                lines.append(f"Notes: {notes}")
    return Extracted("\n".join(lines), f"PowerPoint deck, {_plural(len(deck.slides), 'slide')}")


def _textutil(path: Path) -> Extracted:
    """Older formats, through the converter built into macOS."""
    if sys.platform != "darwin":
        raise DocumentError(
            f"reading {path.suffix.lower()} files needs macOS; convert {path.name} "
            "to .docx first."
        )
    proc = subprocess.run(
        ["textutil", "-convert", "txt", "-stdout", str(path)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=60,
        stdin=subprocess.DEVNULL,
    )
    if proc.returncode != 0:
        raise DocumentError(
            f"could not read {path.name}: {proc.stderr.strip() or 'textutil failed'}"
        )
    return Extracted(proc.stdout, DOCUMENT_KINDS[path.suffix.lower()])


_READERS: dict[str, Callable[[Path], Extracted]] = {
    ".pdf": _pdf,
    ".docx": _docx,
    ".xlsx": _xlsx,
    ".xlsm": _xlsx,
    ".pptx": _pptx,
    ".doc": _textutil,
    ".rtf": _textutil,
    ".odt": _textutil,
}
_PACKAGES = {
    ".pdf": "pypdf",
    ".docx": "python-docx",
    ".xlsx": "openpyxl",
    ".xlsm": "openpyxl",
    ".pptx": "python-pptx",
}


def open_image(path: Path):
    """A picture file as an upright RGB image (phones store rotation in EXIF)."""
    from PIL import Image, ImageOps

    if path.suffix.lower() in (".heic", ".heif"):
        try:
            from pillow_heif import register_heif_opener
        except ImportError as exc:
            raise DocumentError("reading HEIC photos needs the pillow-heif package.") from exc
        register_heif_opener()
    with Image.open(path) as image:
        image.seek(0)  # the first frame of an animation
        image = ImageOps.exif_transpose(image)
        if image.mode in ("RGBA", "LA", "PA") or "transparency" in image.info:
            # Transparent pixels have no colour of their own, and dropping the
            # alpha turns them black - which hides dark text and line art. Show
            # them on white instead.
            layer = image.convert("RGBA")
            backdrop = Image.new("RGBA", layer.size, (255, 255, 255, 255))
            return Image.alpha_composite(backdrop, layer).convert("RGB")
        return image.convert("RGB")


def register(registry: Registry, cfg: Config, llm: LLM) -> None:
    if importlib.util.find_spec("PIL") is None:
        return
    from .screen import fit_for_vision

    @registry.define(
        "view_image",
        description=(
            "Look at a picture file - a photo, screenshot or scan - and answer a "
            "question about it with a vision model. Use it for attached images and "
            "any png, jpg, heic, webp, gif, tiff or bmp file. Ask something specific "
            "when you can."
        ),
        parameters=schema(
            {
                "path": string("The image file."),
                "question": string("What you want to know about it.", default=DEFAULT_QUESTION),
            },
            ["path"],
        ),
    )
    def view_image(path: str, question: str = DEFAULT_QUESTION) -> str:
        target = cfg.resolve_path(path)
        if not target.is_file():
            return f"ERROR: {target} does not exist."
        if target.suffix.lower() not in IMAGE_SUFFIXES:
            return (
                f"ERROR: {target.name} is not a picture view_image can open. "
                "Use read_file for documents."
            )
        try:
            image = open_image(target)
        except DocumentError as exc:
            return f"ERROR: {exc}"
        except Exception as exc:
            return f"ERROR: could not open {target.name}: {exc}"
        answer = llm.look(fit_for_vision(image), question)
        return f"[image {image.width}x{image.height}, {target}]\n{answer}"
