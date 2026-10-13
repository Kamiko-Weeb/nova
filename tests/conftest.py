"""Shared fixtures. Sample files are built fresh in each test's temp folder."""
from __future__ import annotations

import dataclasses
import io
from pathlib import Path

import pytest

from nova.config import Config


@pytest.fixture
def cfg(tmp_path):
    """Nova's config with the workspace pointed at the test's temp folder."""
    return dataclasses.replace(
        Config.from_env(),
        workspace=tmp_path,
        restrict_to_workspace=False,
        log_dir=tmp_path / "logs",
    )


def _pdf_bytes(pages: list[str]) -> bytes:
    """A minimal, valid PDF with one line of Helvetica text per page."""
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "",  # the page tree, filled in once the pages exist
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    kids = []
    for text in pages:
        stream = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET"
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
        objects.append(
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {len(objects)} 0 R /Resources << /Font << /F1 3 0 R >> >> >>"
        )
        kids.append(f"{len(objects)} 0 R")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{offset:010d} 00000 n \n" for offset in offsets).encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref}\n%%EOF\n"
    ).encode()
    return bytes(out)


@pytest.fixture
def make_pdf(tmp_path):
    def make(name: str, pages: list[str], password: str | None = None) -> Path:
        path = tmp_path / name
        path.write_bytes(_pdf_bytes(pages))
        if password is not None:
            from pypdf import PdfReader, PdfWriter

            writer = PdfWriter(clone_from=PdfReader(io.BytesIO(path.read_bytes())))
            writer.encrypt(user_password=password, owner_password="owner", algorithm="RC4-128")
            writer.write(path)
        return path

    return make


@pytest.fixture
def scan_pdf(tmp_path):
    """A PDF that is only a picture, like a scan: no text layer."""
    from PIL import Image

    path = tmp_path / "scan.pdf"
    Image.new("RGB", (200, 100), "white").save(path)
    return path


@pytest.fixture
def sample_docx(tmp_path):
    import docx

    document = docx.Document()
    document.add_paragraph("Intro paragraph")
    table = document.add_table(rows=2, cols=2)
    for (row, col), text in {(0, 0): "Name", (0, 1): "Qty", (1, 0): "Apples", (1, 1): "3"}.items():
        table.cell(row, col).text = text
    document.add_paragraph("Closing paragraph")
    path = tmp_path / "report.docx"
    document.save(path)
    return path


@pytest.fixture
def sample_xlsx(tmp_path):
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Totals"
    sheet.append(["Item", "Price", None])
    sheet.append(["Pen", 2, None])
    sheet.append([None, None, None])
    sheet["B4"] = "=SUM(B2:B2)"  # never calculated, so no value: must not show as text
    workbook.create_sheet("Empty")
    path = tmp_path / "book.xlsx"
    workbook.save(path)
    return path


@pytest.fixture
def sample_pptx(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])  # title only
    slide.shapes.title.text = "Quarterly review"
    group = slide.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1))
    box.text_frame.text = "Grouped note"
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "North"
    slide.notes_slide.notes_text_frame.text = "Mention the north"
    deck.slides.add_slide(deck.slide_layouts[6])  # a blank second slide
    path = tmp_path / "deck.pptx"
    deck.save(path)
    return path


@pytest.fixture(autouse=True)
def _fresh_document_cache():
    from nova.tools import documents

    documents._extract.cache_clear()


@pytest.fixture
def make_image(tmp_path):
    def make(name, size=(64, 32), color="red", exif_orientation=None) -> Path:
        from PIL import Image

        path = tmp_path / name
        if path.suffix.lower() in (".heic", ".heif"):
            from pillow_heif import register_heif_opener

            register_heif_opener()
        image = Image.new("RGB", size, color)
        if exif_orientation:
            exif = Image.Exif()
            exif[0x0112] = exif_orientation  # the EXIF Orientation tag
            image.save(path, exif=exif)
        else:
            image.save(path)
        return path

    return make
